import json
import os
import shutil
import smtplib
from email import encoders
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

from comtypes.client import CreateObject
from docx import Document
from openpyxl import load_workbook
from jira import JIRA

from pattern_receipt import PatternReceipt


class PatternReceiptManager:
    def __init__(self):
        self.load_config()

    def generate_pattern_receipt_word_file(self, pattern_receipt):
        document = Document(self.pattern_receipt_template_path)

        # Replace placeholders in the template with PR data
        for paragraph in document.paragraphs:
            text = paragraph.text
            text = text.replace("{Pattern_Name}", pattern_receipt.pattern_name)
            text = text.replace("{Drawing_Number}", pattern_receipt.drawing_number)
            text = text.replace("{Impressions}", pattern_receipt.impressions)
            text = text.replace("{Pattern_Type}", pattern_receipt.pattern_type)
            text = text.replace("{Core_Box_Boolean}", pattern_receipt.core_box_boolean)
            text = text.replace("{Bore_Box_Name}", pattern_receipt.core_box_name)
            text = text.replace("{Core_Box_Type}", pattern_receipt.core_box_type)
            text = text.replace("{Shipping_Party}", pattern_receipt.shipping_party)
            text = text.replace("{Receiving_Party}", pattern_receipt.receiving_party)
            text = text.replace("{Ship_Date}", pattern_receipt.ship_date)
            text = text.replace("{Notes}", pattern_receipt.notes)
            paragraph.text = text

        word_file_name = f"Pattern Receipt for {pattern_receipt.pattern_name}.docx"
        word_file_path = os.path.join(self.pattern_receipt_Word_folder, word_file_name)
        document.save(word_file_path)

        return word_file_path

    def generate_pattern_receipt_pdf_file(self, pattern_receipt, pattern_receipt_word_file_path):
        word = CreateObject("Word.Application")
        word.Visible = False

        doc = word.Documents.Open(pattern_receipt_word_file_path)
        pdf_file_name = f"Pattern Receipt for {pattern_receipt.pattern_name}.pdf"
        pdf_file_path = os.path.join(self.pattern_receipt_PDF_folder, pdf_file_name)
        doc.SaveAs(pdf_file_path, FileFormat=17)

        doc.Close()
        word.Quit()

    def send_email_to_purchasing(self, pattern_receipt):
        issuer_email = self.personal_email_address

        # Create email message
        msg = MIMEMultipart()
        msg["From"] = issuer_email
        msg["To"] = self.purchasing_email_address
        msg["Subject"] = f"Pattern Receipt for {pattern_receipt.pattern_name}"

        # Attach the receipt PDF
        pattern_receipt_PDF_file_path = os.path.join(self.pattern_receipt_PDF_folder, f"Pattern Receipt for {pattern_receipt.pattern_name}.pdf")
        self.attach_file_to_email(msg, pattern_receipt_PDF_file_path)

        # Add the email body
        body = f"Please see attached Pattern Receipt for {pattern_receipt.pattern_name}."
        msg.attach(MIMEText(body, "plain"))

        # Send the email
        with smtplib.SMTP(self.smtp_server, self.smtp_port) as server:
            server.starttls()
            server.login(self.personal_email_address, self.personal_email_token)
            server.sendmail(issuer_email, self.purchasing_email_address, msg.as_string())

    def attach_file_to_email(self, msg, file_path):
        with open(file_path, "rb") as attachment_file:
            attachment = MIMEBase("application", "octet-stream")
            attachment.set_payload(attachment_file.read())
        encoders.encode_base64(attachment)
        attachment.add_header("Content-Disposition", f"attachment; filename={os.path.basename(file_path)}")
        msg.attach(attachment)

    def attach_pattern_receipt_to_jira_issue(self, pattern_receipt):
        # Define Jira instance details
        project_key = "PT"

        # Authenticate with Jira
        auth = (self.jira_username, self.jira_api_token)
        jira = JIRA(self.jira_base_url, basic_auth=auth)

        # Define issue summary format
        issue_summary = f"{pattern_receipt.pattern_name}+"

        # Find issue with the given summary
        issues = jira.search_issues(f'project="{project}" AND summary ~ "{issue_summary}"', maxResults=1)

        if issues:
            issue = issues[0]
            pdf_filename = f"Pattern Receipt for {pattern_receipt.pattern_name}"

            if os.path.exists(self.pattern_receipt_PDF_folder + "\\" + pdf_filename + ".pdf"):
                # Attach the PDF to the issue
                with open(self.pattern_receipt_PDF_folder + "\\" + pdf_filename + ".pdf", 'rb') as pdf_file:
                    jira.add_attachment(issue, pdf_file, filename=pdf_filename)
                print(f"Successfully attached {pdf_filename} to Jira issue {issue.key}.")
            else:
                print(f"Could not find PDF file '{pdf_filename}'.")
        else:
            print(f"Could not find Jira issue with summary '{issue_summary}' in project '{project_key}'.")

    def create_pattern_receipt(self, receipt_form_user_input):
        pattern_receipt = PatternReceipt.from_user_input(user_input)

        word_file_path = self.generata_pattern_receipt_word_file(pattern_receipt)
        self.generate_pattern_receipt_pdf_file(pattern_receipt, word_file_path)
        self.send_email_to_purchasing(pattern_receipt)
        #self.attach_pattern_receipt_to_jira_issue(pattern_receipt)

        return pattern_receipt

    def load_config(self):
        with open("config.json", "r", encoding="utf-8") as config_file:
            config = json.load(config_file)

        file_paths = config["file_paths"]
        self.pattern_receipt_template_path = file_paths["pattern_receipt_template_file"]
        self.pattern_receipt_Word_folder = file_paths["pattern_receipt_Word_folder"]
        self.patterh_receipt_PDF_folder = file_paths["pattern_receipt_PDF_folder"]

        personal_email_settings = config["personal_email_settings"]
        self.smtp_server = personal_email_settings["smtp_server"]
        self.smtp_port = personal_email_settings["smtp_port"]
        self.personal_email_address = personal_email_settings["personal_email_address"]
        self.personal_email_token = personal_email_settings["personal_email_token"]

        jira_settings = config["jira_settings"]
        self.jira_base_url = jira_settings["jira_base_url"]
        self.jira_username = jira_settings["jira_username"]
        self.jira_api_token = jira_settings["jira_api_token"]

        outgoing_email_addresses = config["outgoing_email_addresses"]
        self.purchasing_email_address = outgoing_email_addresses["purchasing_email_address"]
        
